from django.shortcuts import render
from django.core.files.storage import FileSystemStorage
from django.db.models import Q
from django.http import HttpResponse, HttpResponseRedirect, JsonResponse
from django.urls import reverse, reverse_lazy
from django.views.generic import TemplateView
from django.views.generic.edit import UpdateView, DeleteView
from django_tables2 import SingleTableView
from django.utils import timezone

import xlsxwriter
import openpyxl
import io

from material.models import EstimasiMaterialUsageCO, Product
from .models import ChangeOrder
from .tables import ChangeOrderTable
# Create your views here.

from datetime import date, datetime, timedelta
import decimal

class PPICHomeView(TemplateView):
    template_name = 'ppic/ppic_home.html'

class ChangeOrderListView(SingleTableView):
    model = ChangeOrder
    template_name = 'ppic/change_order_list.html'
    context_table_name = 'change_order_table'
    table_class = ChangeOrderTable

    def get_queryset(self):
        req = self.request.GET.get('co_search_date_input')
        if req:
            query = datetime.strptime(req, '%Y-%m-%d').date()
            return ChangeOrder.objects.filter( Q(date_co=query))
        else:
            return ChangeOrder.objects.filter( Q(date_co=timezone.now().date()))

def daterange(start_date, end_date):
    for n in range(int((end_date - start_date).days)+1):
        yield start_date + timedelta(n)

def count_days(start_date, end_date):
    return 1 + int((end_date - start_date).days)

def import_change_order_xlsx(request):
    if request.method == 'POST':
        request_date = request.POST.get('select_co_date')
        request_file = request.FILES.get('import_change_order_xlsx')
        print('request_co_date = ', request_date, 'type = ', type(request_date))
        print('request_file = ', request_file)
        if request_date and request_file:
            fs = FileSystemStorage()
            file = fs.save(request_file.name, request_file)
            print('file = ', file)
            wb = openpyxl.load_workbook(filename=file, read_only=True)
            ws = wb.active

            change_order_list = []
            emu_list = []
            for row in ws.iter_rows(min_row=2):
                if (row[1].value != None) and (row[3].value != None) and (row[8].value != None):
                    change_order = ChangeOrder()
                    emu_co = EstimasiMaterialUsageCO()
                    change_order.date_co = datetime.strptime(request_date, '%Y-%m-%d').date()
                    change_order.customer = row[0].value
                    change_order.part_id_customer = row[1].value

                    change_order.model = row[2].value
                    change_order.part_name = row[3].value
                    change_order.jumlah_produksi = row[4].value
                    change_order.virgin = row[5].value
                    change_order.regrind = row[6].value
                    change_order.estimasi_selesai = row[7].value
                    change_order.material = row[8].value
                    change_order.cycle_time = row[9].value
                    change_order.cavity = row[10].value
                    change_order.output = row[11].value
                    change_order.no_mesin = row[12].value
                    change_order.prioritas = row[13].value
                    change_order.material_percentage = int(row[14].value * 100)

                    for date in daterange(change_order.date_co, change_order.estimasi_selesai.date()):
                        emu_co.tanggal_co = change_order.date_co
                        emu_co.tanggal_operasi = date
                        emu_co.part_no = change_order.part_id_customer
                        emu_co.target_output = change_order.jumlah_produksi / count_days(change_order.date_co, change_order.estimasi_selesai.date())
                        try:
                            berat_product = Product.objects.filter(part_id = emu_co.part_no).values('part_weight').last()['part_weight']
                        except:
                            berat_product = 0
                        emu_co.berat_target_output = (decimal.Decimal(berat_product) * decimal.Decimal(emu_co.target_output)) /1000
                        if emu_co.tanggal != None and emu_co.part_no != None and emu_co.target_output != None:
                            emu_list.append(emu_co)

                    if change_order.part_id_customer != None and change_order.jumlah_produksi != None:
                        change_order_list.append(change_order)
                        
            try:
                ChangeOrder.objects.bulk_create(change_order_list)
                EstimasiMaterialUsageCO.objects.bulk_create(emu_list)
                wb.close()
                fs.delete(file)

                url = reverse_lazy('change_order_list')
                warning = 'Success importing CO list from Excel file.'

                response = HttpResponse("<script> alert( '%s' ); window.location='%s' </script>" %(warning, url))
            except Exception as e:
                wb.close()
                fs.delete(file)
                url = reverse_lazy('change_order_list')
                print(e)
                response = HttpResponse("<script> alert( '%s' ); window.location='%s' </script>" %(e, url))
    return response

def delete_co(request):
    if request.method == 'POST':
        request_date = request.POST.get('delete_co')
        request_date = datetime.strptime(request_date, '%Y-%m-%d').date()
        print(request_date)
        print(type(request_date))
        if request_date:
            try:
                date_str = request_date.strftime('%d-%m-%Y')
                ChangeOrder.objects.filter(date_co = request_date).delete()
                EstimasiMaterialUsageCO.objects.filter(tanggal_co = request_date).delete()
                url = reverse_lazy('change_order_list')
                return HttpResponse("<script> alert('CO List untuk tanggal %s sudah dihapus'); window.location='%s' </script>" %(date_str, url))
            except Exception as e:
                print(e)
                return HttpResponseRedirect(reverse('change_order_list'))
        else:
            return HttpResponse("<script> alert('Error no request_date');</script>")


def ajax_check_co(request):
    if request.method == 'GET':
        request_date = request.GET.get('select_co_date')
        queryset = list(ChangeOrder.objects.filter(date_co=request_date).values())
        return JsonResponse(data={'query': queryset})